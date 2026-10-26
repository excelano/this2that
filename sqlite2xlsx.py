#!/usr/bin/env python3
# Author: David M. Anderson
# Built with AI assistance (Claude, Anthropic)
"""sqlite2xlsx - Convert a SQLite database to a styled Excel .xlsx workbook.

Usage: sqlite2xlsx <input.db> [output.xlsx] [--table NAME]

Writes one sheet per table, named after the table, each with a bold frozen
header row, an autofilter, and auto-fitted column widths -- the same styling
csv2xlsx applies. Pass --table to write a single-sheet workbook from just
that table. Values keep their SQLite type (integer, real, text); BLOB
columns are written as hex strings, since xlsx has no binary cell type.
"""
import argparse
import re
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Stamped from packaging/control by install.sh and packaging/build-deb.sh.
# A checkout is not stamped, so running from the clone reports the dev
# placeholder rather than claiming a release it is not.
__version__ = "0.0.0+dev"

# Characters openpyxl forbids in a sheet title.
_BAD_TITLE = re.compile(r"[\\/*?:\[\]]")


def fmt(value):
    """Render a SQLite value as an xlsx cell value."""
    if isinstance(value, bytes):
        return value.hex()
    return value


def table_names(conn):
    rows = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
    )
    return [r[0] for r in rows]


def sheet_title(name, used):
    """A unique, Excel-legal sheet title from a table name."""
    title = _BAD_TITLE.sub("_", name)[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def write_sheet(ws, conn, table):
    """Fill one worksheet from one table, with the header/filter/width styling."""
    cur = conn.execute(f'SELECT * FROM "{table}"')
    columns = [d[0] for d in cur.description]
    widths = {c: len(name) for c, name in enumerate(columns, start=1)}

    for c, name in enumerate(columns, start=1):
        ws.cell(row=1, column=c, value=name).font = Font(bold=True)
    for r, row in enumerate(cur, start=2):
        for c, value in enumerate(row, start=1):
            v = fmt(value)
            ws.cell(row=r, column=c, value=v)
            widths[c] = max(widths[c], len(str(v)) if v is not None else 0)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 60)


def main():
    ap = argparse.ArgumentParser(
        prog="sqlite2xlsx", description="Convert a SQLite database to a styled .xlsx workbook."
    )
    ap.add_argument("--version", "-V", action="version", version=f"%(prog)s {__version__}")
    ap.add_argument("input", help="input SQLite database file")
    ap.add_argument("output", nargs="?", help="output .xlsx (default: input name with .xlsx)")
    ap.add_argument("--table", help="export only this table, as a single-sheet workbook")
    args = ap.parse_args()

    try:
        # mode=ro so a bad path surfaces as an error here instead of
        # sqlite3 silently creating a new, empty database file.
        conn = sqlite3.connect(f"file:{args.input}?mode=ro", uri=True)
        conn.execute("SELECT 1")
    except sqlite3.OperationalError as e:
        sys.exit(f"Error: {e}")

    tables = table_names(conn)
    if not tables:
        sys.exit("Error: no tables found")

    if args.table:
        if args.table not in tables:
            sys.exit(f"Error: no table named {args.table!r}. Available: {', '.join(tables)}")
        tables = [args.table]

    wb = Workbook()
    used = set()
    for i, table in enumerate(tables):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sheet_title(table, used)
        write_sheet(ws, conn, table)

    out = args.output or re.sub(r"\.[^.]*$", "", args.input) + ".xlsx"
    wb.save(out)
    print(f"Wrote: {out}")


if __name__ == "__main__":
    main()
