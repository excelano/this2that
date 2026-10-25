#!/usr/bin/env python3
# Author: David M. Anderson
# Built with AI assistance (Claude, Anthropic)
"""ods2xlsx - Convert an OpenDocument .ods spreadsheet to an Excel .xlsx workbook.

Usage: ods2xlsx <input.ods> [output.xlsx]

The reverse of xlsx2ods. Every sheet comes across with typed cells (number,
percentage, date, time, boolean, text), merged cells, named ranges, and
formulas rewritten from OpenFormula into Excel syntax. A formula that cannot
be rewritten (a function Excel has no equivalent for, a reference spanning
sheets) keeps its last computed value and is reported on stderr by cell.
Formulas are written without a saved result; Excel and LibreOffice compute
them on open, but a reader that only looks at saved values (xlsx2csv among
them) sees those cells empty. Sheet names Excel cannot hold are shortened or
cleaned, and formulas follow the new names. Styling, column widths, charts,
and conditional formatting are not carried.
"""
import argparse
import re
import sys
from datetime import date, datetime, timedelta

from odf import teletype
from odf.namespaces import TABLENS
from odf.number import Number, PercentageStyle
from odf.opendocument import load
from odf.style import Style
from odf.table import CalculationSettings, NamedRange, TableCell, TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils import FORMULAE, get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

# Stamped from packaging/control by install.sh and packaging/build-deb.sh.
# A checkout is not stamped, so running from the clone reports the dev
# placeholder rather than claiming a release it is not.
__version__ = "0.0.0+dev"

# OpenFormula names whose Excel counterpart has another name. OpenFormula's
# own CEILING and FLOOR take Excel's .MATH semantics; Excel's CEILING and FLOOR
# arrive as COM.MICROSOFT.CEILING and COM.MICROSOFT.FLOOR.
RENAMED = {
    "FORMULA": "FORMULATEXT",
    "SKEWP": "SKEW.P",
    "CEILING": "CEILING.MATH",
    "FLOOR": "FLOOR.MATH",
}

# Arguments OpenFormula lets a call leave off that Excel requires, by position.
DEFAULTS = {
    "ROUND": {1: "0"},
    "ROUNDUP": {1: "0"},
    "ROUNDDOWN": {1: "0"},
    "NORMDIST": {3: "TRUE"},
    "POISSON": {2: "TRUE"},
    "GAMMADIST": {3: "TRUE"},
    "EXPONDIST": {2: "TRUE"},
    "LOGNORMDIST": {1: "0", 2: "1"},
}

# Functions Excel writes with an _xlfn. prefix that are plain OpenFormula.
NATIVE_XLFN = frozenset("""
ACOT ACOTH ARABIC BASE BITAND BITLSHIFT BITOR BITRSHIFT BITXOR COMBINA COT COTH CSC
CSCH DAYS DECIMAL GAMMA GAUSS IFNA ISFORMULA ISOWEEKNUM MUNIT NUMBERVALUE PDURATION
PERMUTATIONA PHI RRI SEC SECH SHEET SHEETS UNICHAR UNICODE XOR
""".split())

# Dynamic-array functions Excel stores under a second, _xlws. prefix.
XLWS = frozenset({"FILTER", "SORT"})

# Functions whose criteria or patterns follow the document's matching rules.
CRITERIA = re.compile(r"\b(?:COM\.MICROSOFT\.)?(COUNTIFS?|SUMIFS?|AVERAGEIFS?|MAXIFS|MINIFS|MATCH|[HV]LOOKUP|SEARCH"
                      r"|D(?:AVERAGE|COUNTA?|GET|MAX|MIN|PRODUCT|STDEVP?|SUM|VARP?))\(", re.I)

_BAD_TITLE = re.compile(r"[\\/*?:\[\]]")
_REF_PART = re.compile(r"(?:\$?('(?:[^']|'')+'|[^.'\[\]:]+))?\.(\$?[A-Z]+\$?\d+|\$?[A-Z]+|\$?\d+)$")
_TOKEN = re.compile(r"""
    (?P<string>"(?:[^"]|"")*")
  | (?P<ref>\[(?:[^\]']|'(?:[^']|'')*')*\])
  | (?P<error>\#[A-Z0-9/]+[!?]?)
  | (?P<name>[A-Za-z_\\][\w.]*)(?P<call>\s*\()?
  | (?P<number>\d+(?:\.\d*)?(?:[Ee][+-]?\d+)?|\.\d+(?:[Ee][+-]?\d+)?)
  | (?P<sep>;)
  | (?P<row>\|)
  | (?P<union>~)
  | (?P<intersect>!)
  | (?P<other>.)
""", re.X | re.S)


class Untranslatable(Exception):
    pass


def sheet_title(name, used):
    """A unique, Excel-legal sheet title (31 chars, no []:*?/\\)."""
    title = _BAD_TITLE.sub("_", name).strip("'")[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def excel_sheet(name):
    if re.fullmatch(r"[A-Za-z_]\w*", name) and not re.fullmatch(r"[A-Za-z]{1,3}\d+", name):
        return name
    return "'" + name.replace("'", "''") + "'"


def ref_to_excel(body, sheets):
    """An OpenFormula reference body (inside the brackets) as an Excel one."""
    parts = body.split(":")
    if len(parts) > 2:
        raise Untranslatable(f"cannot rewrite reference [{body}]")
    sheet, cells = None, []
    for i, part in enumerate(parts):
        m = _REF_PART.match(part)
        if not m:
            raise Untranslatable(f"cannot rewrite reference [{body}]")
        name = m.group(1)
        if name is not None:
            name = name[1:-1].replace("''", "'") if name.startswith("'") else name
            if i == 0:
                sheet = name
            elif name != sheet:
                raise Untranslatable(f"reference [{body}] spans sheets")
        cells.append(m.group(2))
    if sheet is None:
        return ":".join(cells)
    if sheet not in sheets:
        raise Untranslatable(f"reference [{body}] names an unknown sheet")
    return excel_sheet(sheets[sheet]) + "!" + ":".join(cells)


def function_name(name):
    """Map an OpenFormula function name onto its Excel name, prefix included."""
    upper = name.upper()
    # LibreOffice keeps an Excel function it cannot evaluate under its
    # original _xlfn. name, so Excel can evaluate it again.
    if upper.startswith("_XLFN."):
        return "_xlfn." + upper[6:].replace("_XLWS.", "_xlws.")
    if upper.startswith("COM.MICROSOFT."):
        base = upper[len("COM.MICROSOFT."):]
        return base if base in FORMULAE else xlfn(base)
    if upper.startswith("LEGACY.") and upper[len("LEGACY."):] in FORMULAE:
        return upper[len("LEGACY."):]
    base = RENAMED.get(upper, upper)
    if base in FORMULAE:
        return base
    if base in NATIVE_XLFN or base in RENAMED.values():
        return xlfn(base)
    raise Untranslatable(f"no Excel equivalent for {name}")


def xlfn(name):
    return ("_xlfn._xlws." if name in XLWS else "_xlfn.") + name


def to_excel(formula, sheets, names):
    """Rewrite an OpenFormula formula ("of:=...") as an Excel one ("=...")."""
    _, sep, body = formula.partition(":=")
    if not sep or formula.split(":", 1)[0] not in ("of", "oooc"):
        raise Untranslatable(f"not an OpenFormula formula: {formula}")
    out = []
    # One entry per open bracket: [function name or None, arguments seen, any
    # content in the current argument], so a closing call can fill defaults.
    calls = []
    for m in _TOKEN.finditer(body):
        kind = m.lastgroup if m.lastgroup != "call" else "name"
        text = m.group(0)
        if calls and kind != "sep" and text.strip() and text not in (")", "}"):
            calls[-1][2] = True
        if kind == "ref":
            out.append(ref_to_excel(text[1:-1], sheets))
        elif kind == "name":
            name = m.group("name")
            if m.group("call"):
                excel = function_name(name)
                out.append(excel + "(")
                calls.append([excel, 0, False])
            elif name.upper() in ("TRUE", "FALSE") or name.startswith("_xlpm.") or name in names:
                out.append(name)
            else:
                raise Untranslatable(f"unknown name {name}")
        elif kind == "sep":
            if calls:
                calls[-1][1] += 1
                calls[-1][2] = False
            out.append(",")
        elif kind == "row":
            out.append(";")
        elif kind == "union":
            out.append(",")
        elif kind == "intersect":
            out.append(" ")
        elif text in ("(", "{"):
            calls.append([None, 0, False])
            out.append(text)
        elif text in (")", "}") and calls:
            func, seps, content = calls.pop()
            given = seps + 1 if seps or content else 0
            out.extend("," + v for i, v in sorted(DEFAULTS.get(func, {}).items()) if i >= given)
            out.append(text)
        else:
            out.append(text)
    return "=" + "".join(out)


def iso_value(text):
    """An office:date-value as a date or datetime."""
    if "T" in text:
        return datetime.fromisoformat(text.split(".")[0])
    return date.fromisoformat(text)


def duration(text):
    """An office:time-value (PT13H45M10S) as a timedelta."""
    m = re.fullmatch(r"-?P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:([\d.]+)S)?", text)
    if not m:
        raise ValueError(text)
    d, h, mi, s = m.groups()
    return timedelta(days=int(d or 0), hours=int(h or 0), minutes=int(mi or 0), seconds=float(s or 0))


def number(text):
    v = float(text)
    return int(v) if v.is_integer() and abs(v) < 1e15 else v


def percent_places(doc):
    """Cell style name -> decimal places, for each style showing a percentage."""
    containers = (doc.styles, doc.automaticstyles)
    places = {}
    for container in containers:
        for ps in container.getElementsByType(PercentageStyle):
            n = ps.getElementsByType(Number)
            places[ps.getAttribute("name")] = int(n[0].getAttribute("decimalplaces") or 0) if n else 0
    return {st.getAttribute("name"): places[st.getAttribute("datastylename")]
            for container in containers for st in container.getElementsByType(Style)
            if st.getAttribute("datastylename") in places}


def cell_value(cell, percent):
    """The typed value of a table:table-cell and the number format it needs, or (None, None)."""
    kind = cell.getAttribute("valuetype")
    if kind in ("float", "currency"):
        return number(cell.getAttribute("value")), None
    if kind == "percentage":
        v = float(cell.getAttribute("value"))
        places = percent.get(cell.getAttribute("stylename"), 0 if round(v * 100, 9).is_integer() else 2)
        return v, "0." + "0" * places + "%" if places else "0%"
    if kind == "date":
        v = iso_value(cell.getAttribute("datevalue"))
        if isinstance(v, datetime):
            return v, "yyyy-mm-dd hh:mm:ss"
        return v, "yyyy-mm-dd"
    if kind == "time":
        v = duration(cell.getAttribute("timevalue"))
        if timedelta(0) <= v < timedelta(days=1):
            return (datetime.min + v).time(), "hh:mm:ss"
        return v, "[h]:mm:ss"
    if kind == "boolean":
        return cell.getAttribute("booleanvalue") == "true", None
    text = "\n".join(teletype.extractText(p) for p in cell.getElementsByType(P))
    return (text, None) if kind == "string" or text else (None, None)


def regex_criteria(doc):
    """Whether the document matches criteria as regular expressions, where Excel uses wildcards."""
    found = doc.spreadsheet.getElementsByType(CalculationSettings)
    if not found:
        return True
    settings = found[0]
    return settings.getAttribute("usewildcards") != "true" and settings.getAttribute("useregularexpressions") != "false"


def attr_int(el, name):
    v = el.getAttribute(name)
    return int(v) if v else 1


def write_sheet(ws, table, sheets, names, percent, warn):
    r = 0
    for row in table.getElementsByType(TableRow):
        row_repeat = attr_int(row, "numberrowsrepeated")
        cells = [n for n in row.childNodes
                 if getattr(n, "qname", None) in ((TABLENS, "table-cell"), (TABLENS, "covered-table-cell"))]
        has_content = any(n.qname[1] == "table-cell" and (n.getAttribute("valuetype") or n.getAttribute("formula")
                                                          or n.getElementsByType(P)) for n in cells)
        if not has_content:
            r += row_repeat
            continue
        for _ in range(row_repeat):
            r += 1
            c = 0
            for cell in cells:
                repeat = attr_int(cell, "numbercolumnsrepeated")
                if cell.qname[1] == "covered-table-cell":
                    c += repeat
                    continue
                value, fmt = cell_value(cell, percent)
                formula = cell.getAttribute("formula")
                if value is None and not formula and not cell.getAttribute("numbercolumnsspanned"):
                    c += repeat
                    continue
                for _ in range(repeat):
                    c += 1
                    write_cell(ws, r, c, cell, value, fmt, formula, sheets, names, warn)


def write_cell(ws, r, c, cell, value, fmt, formula, sheets, names, warn):
    target = ws.cell(r, c)
    if formula:
        try:
            text = to_excel(formula, sheets, names)
            cols, rows = attr_int(cell, "numbermatrixcolumnsspanned"), attr_int(cell, "numbermatrixrowsspanned")
            if cell.getAttribute("numbermatrixcolumnsspanned"):
                ref = f"{target.coordinate}:{get_column_letter(c + cols - 1)}{r + rows - 1}"
                value = ArrayFormula(ref, text)
            else:
                value = text
        except Untranslatable as e:
            warn(f"{ws.title}!{target.coordinate}: kept the value; {e}")
    if value is not None:
        target.value = value
        if fmt:
            target.number_format = fmt
    cols, rows = attr_int(cell, "numbercolumnsspanned"), attr_int(cell, "numberrowsspanned")
    if cols > 1 or rows > 1:
        ws.merge_cells(start_row=r, start_column=c, end_row=r + rows - 1, end_column=c + cols - 1)


def main():
    ap = argparse.ArgumentParser(
        prog="ods2xlsx", description="Convert an OpenDocument .ods spreadsheet to an Excel .xlsx workbook."
    )
    ap.add_argument("--version", "-V", action="version", version=f"%(prog)s {__version__}")
    ap.add_argument("input", help="input .ods file")
    ap.add_argument("output", nargs="?", help="output .xlsx (default: input name with .xlsx)")
    args = ap.parse_args()

    try:
        doc = load(args.input)
    except Exception as e:
        sys.exit(f"Error: {e}")
    if doc.mimetype != "application/vnd.oasis.opendocument.spreadsheet":
        sys.exit(f"Error: not an OpenDocument spreadsheet: {args.input}")

    tables = [t for t in doc.spreadsheet.childNodes if getattr(t, "qname", None) == (TABLENS, "table")]
    if not tables:
        sys.exit("Error: no sheets found")
    used = set()
    sheets = {t.getAttribute("name"): sheet_title(t.getAttribute("name"), used) for t in tables}

    warnings = []
    warn = warnings.append
    names = {}
    for nr in doc.spreadsheet.getElementsByType(NamedRange):
        name, address = nr.getAttribute("name"), nr.getAttribute("cellrangeaddress")
        try:
            names[name] = ref_to_excel(address, sheets)
        except Untranslatable as e:
            warn(f"name {name}: not converted; {e}")

    wb = Workbook()
    wb.remove(wb.active)
    percent = percent_places(doc)
    for table in tables:
        write_sheet(wb.create_sheet(sheets[table.getAttribute("name")]), table, sheets, names, percent, warn)
    for name, ref in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    if regex_criteria(doc) and any(CRITERIA.search(c.getAttribute("formula") or "")
                                   for c in doc.spreadsheet.getElementsByType(TableCell)):
        warn("this file matches criteria as regular expressions and Excel matches them as wildcards, "
             "so COUNTIF, SUMIF, MATCH, and lookups given a pattern may give different results")

    out = args.output or re.sub(r"\.[^.]*$", "", args.input) + ".xlsx"
    wb.save(out)
    for w in warnings:
        print(f"Warning: {w}", file=sys.stderr)
    print(f"Wrote: {out}")


if __name__ == "__main__":
    main()
